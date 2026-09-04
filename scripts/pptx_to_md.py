#!/usr/bin/env python3
"""PPT 稿 → Markdown：方案评审 PPT、立项汇报 PPT 转成 agent 能整段处理的 md，嵌入图抽到旁边的 _media 文件夹。

Copyright (c) 2026 handsomestWei
改编自 handsomestWei/patent-disclosure-skill（MIT）tools/pptx_to_md.py，改动：每页标题并入 "## 第 N 页：标题"；
表格补 Markdown 表头分隔行；命令行对齐本包 docx_to_md.py（位置参数 + -o，也兼容 -i）；缺 python-pptx 退出码 2；
输出统一 OK:/FAIL: 机读行；主体拆成可 import 的 convert()。

用法：
  python3 pptx_to_md.py 评审.pptx                 # 输出 评审.md + 评审_media/
  python3 pptx_to_md.py 评审.pptx -o 输出.md
  python3 pptx_to_md.py -i 评审.pptx -o 输出.md   # 与源项目一致的写法也认

依赖：pip3 install python-pptx
每页一个 "## 第 N 页：标题" 小节，下面依次是正文文本框、表格（转成 Markdown 表）、嵌入图片引用、"**备注**：" 演讲者备注。
旧版 .ppt（二进制）不支持，请先在 PowerPoint / WPS 里另存为 .pptx。

成功时 stdout 最后一行：OK: slides=.. paragraphs=.. tables=.. images=.. notes=.. out=.. media_dir=..
失败时 stderr 一行 FAIL: 原因；退出码 2 = 缺依赖，1 = 输入或运行错误。
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path


class _Parser(argparse.ArgumentParser):
    """参数错误也走 FAIL: 行，退出码 1（argparse 默认 2，会和"缺依赖"混淆）。"""

    def error(self, message):
        print(f"FAIL: 参数错误：{message}", file=sys.stderr)
        raise SystemExit(1)


class DependencyMissing(SystemExit):
    """缺依赖：退出码固定 2，str() 是人话。"""

    def __init__(self, msg: str):
        super().__init__(2)
        self.msg = msg

    def __str__(self) -> str:
        return self.msg


def _require_pptx():
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        msg = "需要先装 python-pptx：pip3 install python-pptx"
        print(f"FAIL: {msg}", file=sys.stderr)
        raise DependencyMissing(msg)
    return Presentation, MSO_SHAPE_TYPE


def _walk_shapes(shapes, MSO_SHAPE_TYPE):
    """组合形状拆开，按出现顺序逐个给出。"""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes, MSO_SHAPE_TYPE)
        else:
            yield shape


def _text_lines(shape) -> list[str]:
    """文本框按段落取，空段落去掉；带项目符号层级的段落前面加 "- "。"""
    out = []
    for para in shape.text_frame.paragraphs:
        text = "".join(run.text for run in para.runs).strip() or (para.text or "").strip()
        if not text:
            continue
        out.append(("- " if para.level > 0 else "") + text)
    return out


def _table_lines(shape) -> list[str]:
    rows = []
    for row in shape.table.rows:
        cells = [(cell.text or "").strip().replace("\n", " ").replace("|", "\\|") for cell in row.cells]
        rows.append(cells)
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    for r in rows[1:]:
        lines.append("| " + " | ".join(r) + " |")
    return lines


def convert(pptx_path: Path, out_path: Path, media_dir: Path | None = None) -> dict:
    Presentation, MSO_SHAPE_TYPE = _require_pptx()
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    media_dir = Path(media_dir) if media_dir else out_path.parent / f"{out_path.stem}_media"

    prs = Presentation(str(pptx_path))
    stats = {"slides": 0, "paragraphs": 0, "tables": 0, "images": 0, "notes": 0, "warnings": []}
    lines: list[str] = [f"<!-- 由 pptx_to_md.py 自 {Path(pptx_path).name} 转换 -->", ""]
    img_no = 0

    for sn, slide in enumerate(prs.slides, start=1):
        stats["slides"] += 1
        title_shape = slide.shapes.title
        title = (title_shape.text_frame.text or "").strip().replace("\n", " ") if title_shape is not None else ""
        lines.append(f"## 第 {sn} 页：{title}" if title else f"## 第 {sn} 页")
        lines.append("")

        for shape in _walk_shapes(slide.shapes, MSO_SHAPE_TYPE):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    img = shape.image
                    ext = (img.ext or "png").lower()
                    ext = "jpg" if ext == "jpeg" else ext
                    img_no += 1
                    media_dir.mkdir(parents=True, exist_ok=True)
                    fname = f"slide{sn:02d}_img{img_no:03d}.{ext}"
                    (media_dir / fname).write_bytes(img.blob)
                    lines.append(f"![]({media_dir.name}/{fname})")
                    lines.append("")
                    stats["images"] += 1
                except Exception as e:  # noqa: BLE001
                    stats["warnings"].append(f"第 {sn} 页图片抽取失败：{e}")
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                tl = _table_lines(shape)
                if tl:
                    lines.extend(tl)
                    lines.append("")
                    stats["tables"] += 1
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                tl = _text_lines(shape)
                if tl:
                    lines.extend(tl)
                    lines.append("")
                    stats["paragraphs"] += len(tl)

        try:
            if slide.has_notes_slide:
                note = (slide.notes_slide.notes_text_frame.text or "").strip()
                if note:
                    lines.append("**备注**：")
                    lines.append("")
                    lines.extend(ln.strip() for ln in note.splitlines() if ln.strip())
                    lines.append("")
                    stats["notes"] += 1
        except (AttributeError, ValueError) as e:
            stats["warnings"].append(f"第 {sn} 页备注读取失败：{e}")

    out_path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    stats["output"] = str(out_path)
    stats["media_dir"] = str(media_dir) if stats["images"] else None
    return stats


def main(argv=None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = _Parser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx", nargs="?", help="输入 .pptx / .ppsx")
    ap.add_argument("-i", "--input", default=None, help="同位置参数，兼容 -i 写法")
    ap.add_argument("-o", "--output", default=None)
    ap.add_argument("--media-dir", default=None, help="图片输出目录（默认与 .md 同级的 {主名}_media）")
    args = ap.parse_args(argv)
    src_arg = args.pptx or args.input
    if not src_arg:
        print("FAIL: 没给输入文件（位置参数或 -i）", file=sys.stderr)
        return 1
    src = Path(src_arg)
    if not src.is_file():
        print(f"FAIL: 找不到 {src}", file=sys.stderr)
        return 1
    if src.suffix.lower() not in (".pptx", ".ppsx"):
        print(f"FAIL: 只支持 .pptx / .ppsx（旧版 .ppt 请先另存为 .pptx）：{src.name}", file=sys.stderr)
        return 1
    out = Path(args.output) if args.output else src.with_suffix(".md")
    try:
        st = convert(src, out, Path(args.media_dir) if args.media_dir else None)
    except SystemExit:
        raise
    except Exception as e:  # noqa: BLE001
        print(f"FAIL: 演示文稿打不开或转换出错：{e}", file=sys.stderr)
        return 1
    print(f"{st['slides']} 页、{st['paragraphs']} 段、{st['tables']} 张表、{st['images']} 张图、"
          f"{st['notes']} 页有备注 → {out}")
    for w in st["warnings"]:
        print("警告：" + w, file=sys.stderr)
    print(f"OK: slides={st['slides']} paragraphs={st['paragraphs']} tables={st['tables']} "
          f"images={st['images']} notes={st['notes']} out={out} media_dir={st['media_dir'] or '-'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

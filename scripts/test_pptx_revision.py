"""pptx_to_md.py 与 revision_log.py 的单元测试。pptx 部分需要 python-pptx（见 requirements-full.txt），没装就跳过。"""

from __future__ import annotations

import contextlib
import io
import struct
import sys
import tempfile
import unittest
import zlib
from datetime import datetime, timedelta, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import pptx_to_md  # noqa: E402
import revision_log  # noqa: E402


def 有(模块: str) -> bool:
    import importlib.util
    return importlib.util.find_spec(模块) is not None


需要pptx = unittest.skipUnless(有("pptx"), "没装 python-pptx（pip install python-pptx）")


def 跑main(func, argv):
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        try:
            rc = func(argv)
        except SystemExit as e:
            rc = e.code
    return rc, out.getvalue(), err.getvalue()


def 最后一行(text: str) -> str:
    lines = [ln for ln in text.splitlines() if ln.strip()]
    return lines[-1] if lines else ""


def 小PNG(w: int = 4, h: int = 4) -> bytes:
    """不靠 matplotlib，手工拼一张合法的白色 PNG。"""
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    raw = b"".join(b"\x00" + b"\xff" * (w * 3) for _ in range(h))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def 造PPT(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    # 第 1 页：标题页
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "调度方案评审"
    s1.placeholders[1].text = "副标题一行"
    # 第 2 页：标题 + 多级要点 + 备注
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "痛点"
    tf = s2.placeholders[1].text_frame
    tf.text = "资源错配"
    p = tf.add_paragraph()
    p.text = "静态优先级"
    p.level = 1
    s2.notes_slide.notes_text_frame.text = "讲到这里停 10 秒\n强调限频"
    # 第 3 页：空白版式 + 表格 + 图片 + 组合形状里的文本框
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    tbl = s3.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text = "参数"
    tbl.cell(0, 1).text = "取值"
    tbl.cell(1, 0).text = "窗口 W"
    tbl.cell(1, 1).text = "8 | 16"
    png = path.parent / "pic.png"
    png.write_bytes(小PNG())
    s3.shapes.add_picture(str(png), Inches(1), Inches(3), Inches(1), Inches(1))
    grp = s3.shapes.add_group_shape()
    tb = grp.shapes.add_textbox(Inches(3), Inches(3), Inches(2), Inches(1))
    tb.text_frame.text = "组合里的说明"
    prs.save(str(path))


@需要pptx
class PptxTest(unittest.TestCase):
    def test_转换内容齐全(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            造PPT(td / "deck.pptx")
            st = pptx_to_md.convert(td / "deck.pptx", td / "out" / "deck.md")
            md = (td / "out" / "deck.md").read_text(encoding="utf-8")
            self.assertEqual(st["slides"], 3)
            self.assertEqual(st["tables"], 1)
            self.assertEqual(st["images"], 1)
            self.assertEqual(st["notes"], 1)
            self.assertIn("## 第 1 页：调度方案评审", md)
            self.assertIn("## 第 2 页：痛点", md)
            self.assertIn("## 第 3 页\n", md)                       # 空白版式没标题
            self.assertIn("副标题一行", md)
            self.assertIn("资源错配\n- 静态优先级", md)              # 二级要点带 "- "
            self.assertEqual(md.count("痛点"), 1)                    # 标题不重复出现在正文
            self.assertIn("| 参数 | 取值 |\n|---|---|\n| 窗口 W | 8 \\| 16 |", md)
            self.assertIn("![](deck_media/slide03_img001.png)", md)
            self.assertIn("**备注**：\n\n讲到这里停 10 秒\n强调限频", md)
            self.assertIn("组合里的说明", md)
            media = list((td / "out" / "deck_media").glob("*.png"))
            self.assertEqual(len(media), 1)
            self.assertTrue(media[0].read_bytes().startswith(b"\x89PNG"))
            self.assertEqual(st["media_dir"], str(td / "out" / "deck_media"))

    def test_main_两种命令行写法与OK行(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            造PPT(td / "deck.pptx")
            rc, out, err = 跑main(pptx_to_md.main, [str(td / "deck.pptx")])
            self.assertEqual(rc, 0, err)
            self.assertTrue((td / "deck.md").is_file())
            last = 最后一行(out)
            self.assertTrue(last.startswith("OK: slides=3 "), out)
            self.assertIn("images=1", last)
            self.assertIn("media_dir=", last)
            rc, out, _ = 跑main(pptx_to_md.main, ["-i", str(td / "deck.pptx"), "-o", str(td / "b" / "x.md"),
                                                 "--media-dir", str(td / "b" / "imgs")])
            self.assertEqual(rc, 0)
            self.assertIn("![](imgs/slide03_img001.png)", (td / "b" / "x.md").read_text(encoding="utf-8"))
            self.assertTrue((td / "b" / "imgs" / "slide03_img001.png").is_file())

    def test_输入错误FAIL退出1(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            rc, _, err = 跑main(pptx_to_md.main, [str(td / "没有.pptx")])
            self.assertEqual(rc, 1)
            self.assertTrue(err.startswith("FAIL: "), err)
            (td / "old.ppt").write_bytes(b"\xd0\xcf\x11\xe0")
            rc, _, err = 跑main(pptx_to_md.main, [str(td / "old.ppt")])
            self.assertEqual(rc, 1)
            self.assertIn(".ppt", err)
            rc, _, err = 跑main(pptx_to_md.main, [])
            self.assertEqual(rc, 1)
            self.assertIn("FAIL", err)

    def test_没装python_pptx时退出2(self):
        saved = {k: sys.modules.get(k) for k in list(sys.modules) if k == "pptx" or k.startswith("pptx.")}
        for k in saved:
            sys.modules[k] = None                     # 让 from pptx import ... 抛 ImportError
        sys.modules["pptx"] = None
        try:
            with tempfile.TemporaryDirectory() as td:
                (Path(td) / "a.pptx").write_bytes(b"PK")
                rc, _, err = 跑main(pptx_to_md.main, [str(Path(td) / "a.pptx")])
                self.assertEqual(rc, 2)
                self.assertIn("FAIL: 需要先装 python-pptx", err)
        finally:
            sys.modules.pop("pptx", None)
            for k, v in saved.items():
                sys.modules[k] = v


class RevisionLogTest(unittest.TestCase):
    北京 = timezone(timedelta(hours=8))

    def test_新建带表头再追加序号顺延(self):
        with tempfile.TemporaryDirectory() as td:
            case = Path(td)
            t1 = datetime(2026, 9, 4, 11, 22, 33, tzinfo=self.北京)
            r1 = revision_log.append_entry(case, "改稿", "把权 1 的限频改成可配置",
                                           ["权利要求书_v2.md", "说明书_v2.docx"], "权 1 加'预设间隔'限定", now=t1)
            self.assertEqual(r1["entry"], 1)
            self.assertTrue(r1["created"])
            log = case / "修订记录.md"
            text = log.read_text(encoding="utf-8")
            self.assertTrue(text.startswith("# 修订记录\n"))
            self.assertIn("revision_log.py", text)                          # 表头说明
            self.assertIn("## 1 · 2026-09-04 11:22 · 改稿", text)
            self.assertIn("2026-09-04 11:22:33 +0800（UTC 2026-09-04 03:22:33）", text)  # 本地 + UTC
            self.assertIn("- 类型：改稿", text)
            self.assertIn("- 用户这轮说的：把权 1 的限频改成可配置", text)
            self.assertIn("- 交付文件：`权利要求书_v2.md`、`说明书_v2.docx`", text)
            self.assertIn("- 改了什么、为什么：权 1 加'预设间隔'限定", text)

            r2 = revision_log.append_entry(case, "检索", "再查 D3\n要快", [], "补检索 CN1234，\n没发现限频",
                                           now=t1 + timedelta(hours=1))
            self.assertEqual(r2["entry"], 2)
            self.assertFalse(r2["created"])
            text = log.read_text(encoding="utf-8")
            self.assertEqual(text.count("# 修订记录\n"), 1)                 # 表头只写一次
            self.assertIn("## 2 · 2026-09-04 12:22 · 检索", text)
            self.assertIn("- 交付文件：（无新文件）", text)
            self.assertIn("- 用户这轮说的：再查 D3 要快", text)                # 换行压成空格
            self.assertIn("- 改了什么、为什么：补检索 CN1234， 没发现限频", text)
            self.assertLess(text.index("## 1 ·"), text.index("## 2 ·"))

    def test_类型校验与目录不存在(self):
        with tempfile.TemporaryDirectory() as td:
            with self.assertRaises(ValueError):
                revision_log.append_entry(Path(td), "乱写", "u", [], "s")
            with self.assertRaises(FileNotFoundError):
                revision_log.append_entry(Path(td) / "没有", "改稿", "u", [], "s")
            rc, _, err = 跑main(revision_log.main, ["--case-dir", str(Path(td) / "没有"), "--kind", "改稿",
                                                    "--user", "u", "--summary", "s"])
            self.assertEqual(rc, 1)
            self.assertTrue(err.startswith("FAIL: "), err)
            rc, _, err = 跑main(revision_log.main, ["--case-dir", td, "--kind", "乱写",
                                                    "--user", "u", "--summary", "s"])
            self.assertEqual(rc, 1)
            self.assertIn("FAIL: 参数错误", err)

    def test_main_OK行(self):
        with tempfile.TemporaryDirectory() as td:
            rc, out, err = 跑main(revision_log.main, ["--case-dir", td, "--kind", "答复", "--user", "按 OA1 改",
                                                      "--files", "答复意见.md, 修改对照表.md", "--summary", "区别特征改写"])
            self.assertEqual(rc, 0, err)
            self.assertEqual(最后一行(out), f"OK: entry=1 kind=答复 files=2 out={Path(td) / '修订记录.md'}")
            self.assertIn("`答复意见.md`、`修改对照表.md`", (Path(td) / "修订记录.md").read_text(encoding="utf-8"))
            rc, out, _ = 跑main(revision_log.main, ["--case-dir", td, "--kind", "其他", "--user", "x", "--summary", "y"])
            self.assertTrue(最后一行(out).startswith("OK: entry=2 kind=其他 files=0"))


if __name__ == "__main__":
    unittest.main()

"""转换与附图脚本的单元测试。需要 python-docx / pypdf / matplotlib（见 requirements-min.txt）。"""

from __future__ import annotations  # 3.9 上 `str | None` 注解要靠它，包声明支持 3.9+

import contextlib
import io
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import check_env  # noqa: E402
import docx_to_md  # noqa: E402
import make_figure  # noqa: E402
import md_to_docx  # noqa: E402
import pdf_to_text  # noqa: E402
import math_to_omml  # noqa: E402


def 有(模块: str) -> bool:
    """可选依赖在不在。没装就跳过对应用例，而不是报一片 ERROR——
    买家不装 requirements-min.txt 直接跑测试时，看到 11 个红字会以为包坏了。"""
    import importlib.util
    return importlib.util.find_spec(模块) is not None


需要docx = unittest.skipUnless(有("docx"), "没装 python-docx（pip install -r requirements-min.txt）")
需要pypdf = unittest.skipUnless(有("pypdf"), "没装 pypdf（pip install -r requirements-min.txt）")
需要matplotlib = unittest.skipUnless(有("matplotlib"), "没装 matplotlib（pip install -r requirements-min.txt）")
需要latex2mathml = unittest.skipUnless(有("latex2mathml"), "没装 latex2mathml（pip install -r requirements-min.txt）")


def 跑main(func, argv):
    """跑某个脚本的 main()，返回 (退出码, stdout, stderr)。"""
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        try:
            rc = func(argv)
        except SystemExit as e:
            rc = e.code
    return rc, out.getvalue(), err.getvalue()


def 最后一行(text: str) -> str:
    lines = [ln for ln in text.splitlines() if ln.strip()]
    return lines[-1] if lines else ""

MD_SAMPLE = """# 发明名称：一种演练用的方法

## 技术领域

本方法属于测试领域。

## 具体实施方式

步骤一，做第一件事。

![](图1.png)

| 参数 | 取值 |
|---|---|
| 窗口 | 120 秒 |
| 倍数 | 2 |

1. 一种方法，其特征在于，包括步骤A。
"""

FIG_SPEC = {
    "type": "flow", "title": "图1 测试流程",
    "nodes": [
        {"id": "S101", "label": "接收消息"},
        {"id": "S102", "label": "判断窗口", "shape": "diamond"},
        {"id": "S103", "label": "发出告警", "shape": "rounded"},
    ],
    "edges": [["S101", "S102"], ["S102", "S103", "窗口外"]],
}


@需要matplotlib
class FigureTest(unittest.TestCase):
    def test_流程图出PNG(self):
        with tempfile.TemporaryDirectory() as td:
            spec = Path(td) / "图1.json"
            spec.write_text(json.dumps(FIG_SPEC, ensure_ascii=False), encoding="utf-8")
            out = Path(td) / "图1.png"
            r = make_figure.draw(json.loads(spec.read_text(encoding="utf-8")), out)
            self.assertTrue(out.is_file() and out.stat().st_size > 5000)
            self.assertEqual(r["nodes"], 3)
            self.assertEqual(r["edges"], 2)
            self.assertTrue(out.read_bytes()[:8].startswith(b"\x89PNG"))

    def test_框图分行布局(self):
        spec = {"type": "block", "nodes": [{"id": "101", "label": "模块甲"},
                                           {"id": "102", "label": "模块乙"},
                                           {"id": "103", "label": "模块丙"}],
                "edges": [["101", "102"], ["101", "103"]],
                "rows": [["101"], ["102", "103"]]}
        pos = make_figure.layout(spec)
        self.assertEqual(pos["101"][1], 0)
        self.assertEqual(pos["102"][1], pos["103"][1])
        self.assertLess(pos["102"][0], pos["103"][0])

    def test_画图不改全局字体(self):
        """draw() 只在自己这次调用里换中文字体。

        local_tools.py 的 MCP 进程长驻：这里漏一次，后面所有 matplotlib 输出都跟着换，
        连带把同进程生成的 PDF 变成抽不出文字的那种。
        """
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        before = list(plt.rcParams["font.family"])
        with tempfile.TemporaryDirectory() as td:
            make_figure.draw(FIG_SPEC, Path(td) / "f.png")
        self.assertEqual(list(plt.rcParams["font.family"]), before)

    def test_main_端到端(self):
        with tempfile.TemporaryDirectory() as td:
            spec = Path(td) / "s.json"
            spec.write_text(json.dumps(FIG_SPEC), encoding="utf-8")
            rc = make_figure.main([str(spec)])
            self.assertEqual(rc, 0)
            self.assertTrue((Path(td) / "s.png").is_file())


@需要docx
class RoundTripTest(unittest.TestCase):
    def test_md_到docx_再回md(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            # 先画一张真图作为附图
            fig = td / "图1.png"
            make_figure.draw(FIG_SPEC, fig)
            (td / "稿.md").write_text(MD_SAMPLE, encoding="utf-8")

            st = md_to_docx.convert(td / "稿.md", td / "稿.docx")
            self.assertEqual(st["headings"], 3)
            self.assertEqual(st["images"], 1)
            self.assertEqual(st["tables"], 1)
            self.assertEqual(st["missing_images"], [])
            self.assertTrue((td / "稿.docx").stat().st_size > 10000)

            st2 = docx_to_md.convert(td / "稿.docx", td / "回转.md")
            back = (td / "回转.md").read_text(encoding="utf-8")
            self.assertEqual(st2["headings"], 3)
            self.assertIn("# 发明名称", back)
            self.assertEqual(st2["images"], 1)
            self.assertEqual(st2["tables"], 1)
            self.assertIn("发明名称：一种演练用的方法", back)
            self.assertIn("步骤一，做第一件事。", back)
            self.assertIn("| 窗口 | 120 秒 |", back)
            self.assertIn("![](回转_media/", back)
            media = list((td / "回转_media").glob("*"))
            self.assertEqual(len(media), 1)
            self.assertTrue(media[0].read_bytes()[:8].startswith(b"\x89PNG"))

    def test_权利要求编号不被当成列表(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            (td / "c.md").write_text("## 权利要求书\n\n1. 一种方法，其特征在于X。\n", encoding="utf-8")
            md_to_docx.convert(td / "c.md", td / "c.docx")
            docx_to_md.convert(td / "c.docx", td / "c2.md")
            self.assertIn("1. 一种方法", (td / "c2.md").read_text(encoding="utf-8"))

    def test_缺图给出提示并返回1(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            (td / "m.md").write_text("段落\n\n![](不存在.png)\n", encoding="utf-8")
            rc = md_to_docx.main([str(td / "m.md")])
            self.assertEqual(rc, 1)
            self.assertTrue((td / "m.docx").is_file())


def _造png(path: Path, w_px: int, h_px: int):
    """造一张指定像素尺寸的 PNG，用来验证按宽高比定尺寸的逻辑。"""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig = plt.figure(figsize=(w_px / 100, h_px / 100), dpi=100)
    fig.savefig(str(path), dpi=100)
    plt.close(fig)


@需要docx
class MarkdownQuirksTest(unittest.TestCase):
    """mermaid 出图那一路带出来的三个坑：HTML 注释、长图截断、``` 代码块。"""

    def test_整行HTML注释不进正文(self):
        """mermaid_render 会在 md 里留 <!-- ![图示 1](fig_001.png) --> 这种行，
        以前会原样当正文写进 Word。"""
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            (td / "a.md").write_text(
                "正文一。\n\n"
                "<!-- mermaid source: fig_001.mmd -->\n"
                "<!-- ![图示 1](mermaid_figures/fig_001.png) -->\n\n"
                "正文二。\n",
                encoding="utf-8")
            md_to_docx.convert(td / "a.md", td / "a.docx")
            import docx as _d
            texts = [p.text for p in _d.Document(str(td / "a.docx")).paragraphs]
            self.assertNotIn(True, ["<!--" in t for t in texts], texts)
            self.assertIn("正文一。", texts)
            self.assertIn("正文二。", texts)

    def test_多行HTML注释块整块跳过(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            (td / "b.md").write_text(
                "正文。\n\n<!--\n藏起来的说明\n还有一行\n-->\n\n尾巴。\n", encoding="utf-8")
            md_to_docx.convert(td / "b.md", td / "b.docx")
            import docx as _d
            texts = [p.text for p in _d.Document(str(td / "b.docx")).paragraphs]
            self.assertNotIn("藏起来的说明", texts)
            self.assertNotIn("还有一行", texts)
            self.assertIn("尾巴。", texts)

    @需要matplotlib
    def test_竖长图限高不被截(self):
        """1298×4084 那种长流程图，铺满 14.6cm 宽会有 45.9cm 高，超出 A4 版心直接被截。"""
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            高图, 正常图 = td / "tall.png", td / "wide.png"
            _造png(高图, 130, 409)          # 和 1298×4084 同一个宽高比
            _造png(正常图, 400, 300)
            self.assertEqual(list(md_to_docx.picture_size(高图)), ["height"])
            self.assertEqual(list(md_to_docx.picture_size(正常图)), ["width"])

            (td / "c.md").write_text("![](tall.png)\n\n![](wide.png)\n", encoding="utf-8")
            st = md_to_docx.convert(td / "c.md", td / "c.docx")
            self.assertEqual(st["images"], 2)
            import docx as _d
            from docx.shared import Cm
            形状 = _d.Document(str(td / "c.docx")).inline_shapes
            self.assertEqual(len(形状), 2)
            for sh in 形状:                # 两张都要落在版心内
                self.assertLessEqual(sh.height, Cm(md_to_docx.MAX_IMG_H_CM) + Cm(0.1))
                self.assertLessEqual(sh.width, Cm(md_to_docx.MAX_IMG_CM) + Cm(0.1))
            self.assertAlmostEqual(形状[0].height.cm, md_to_docx.MAX_IMG_H_CM, places=1)
            self.assertAlmostEqual(形状[1].width.cm, md_to_docx.MAX_IMG_CM, places=1)

    def test_读不出尺寸时退回按宽度(self):
        with tempfile.TemporaryDirectory() as td:
            假图 = Path(td) / "x.png"
            假图.write_bytes(b"not a real png")
            self.assertEqual(list(md_to_docx.picture_size(假图)), ["width"])

    def test_代码块用等宽字体且不印围栏行(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            (td / "d.md").write_text(
                "前言。\n\n"
                "```mermaid\n"
                "flowchart TB\n"
                "  A[\"起点\"] --> B[\"终点\"]\n"
                "```\n\n"
                "后记。\n",
                encoding="utf-8")
            md_to_docx.convert(td / "d.md", td / "d.docx")
            import docx as _d
            段 = _d.Document(str(td / "d.docx")).paragraphs
            texts = [p.text for p in 段]
            self.assertNotIn(True, [t.strip().startswith("```") for t in texts], texts)
            self.assertIn("flowchart TB", texts)
            self.assertIn("前言。", texts)
            self.assertIn("后记。", texts)
            等宽 = [p.text for p in 段 if any(r.font.name == md_to_docx.MONO_FONT_EN for r in p.runs)]
            self.assertIn("flowchart TB", 等宽)
            self.assertNotIn("前言。", 等宽)

    def test_代码块里的井号和公式不被解析(self):
        """围栏里是原样代码：# 不是标题，$...$ 也不该被转成公式。"""
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            (td / "e.md").write_text(
                "## 真标题\n\n```python\n# 这是注释不是标题\nx = 1  # $a^2$ 也不是公式\n```\n",
                encoding="utf-8")
            st = md_to_docx.convert(td / "e.md", td / "e.docx")
            self.assertEqual(st["headings"], 1)
            self.assertEqual(st["omml"], 0)
            import docx as _d
            texts = [p.text for p in _d.Document(str(td / "e.docx")).paragraphs]
            self.assertIn("# 这是注释不是标题", texts)
            self.assertIn("x = 1  # $a^2$ 也不是公式", texts)

    def test_三个坑一起出现时正文表格公式都还在(self):
        """回归：修坑不能把原来的标题/表格/公式/字体弄丢。"""
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            (td / "f.md").write_text(
                "# 标题\n\n"
                "行内公式 \\(x_1\\) 在这里。\n\n"
                "```mermaid\nflowchart TD\n  A --> B\n```\n"
                "<!-- ![图示 1](fig_001.png) -->\n\n"
                "| 参数 | 取值 |\n| --- | --- |\n| 窗口 | 120 |\n",
                encoding="utf-8")
            st = md_to_docx.convert(td / "f.md", td / "f.docx")
            self.assertEqual(st["headings"], 1)
            self.assertEqual(st["tables"], 1)
            import docx as _d
            d = _d.Document(str(td / "f.docx"))
            self.assertEqual(len(d.tables), 1)
            self.assertEqual(d.tables[0].rows[1].cells[0].text, "窗口")
            texts = [p.text for p in d.paragraphs]
            self.assertNotIn(True, ["<!--" in t for t in texts], texts)
            正文 = [p for p in d.paragraphs if p.text.startswith("行内公式")]
            self.assertTrue(正文)
            self.assertEqual(正文[0].runs[0].font.name, md_to_docx.BODY_FONT_EN)


@需要pypdf
class PdfTest(unittest.TestCase):
    def _make_pdf(self, path: Path, text: str | None, font: str = "DejaVu Sans"):
        """造一份测试 PDF。

        字体必须显式指定：默认的 DejaVu Sans 带 ToUnicode 表，代表"正常带文字层的
        对比文件"；换成中文字体则代表"有字符但抽不出文字"的那一类 PDF。
        不写死字体的话，同进程里别的用例改过 rcParams 就会串味。
        """
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        with plt.rc_context({"font.family": [font]}):
            fig, ax = plt.subplots(figsize=(8.27, 11.69))
            ax.axis("off")
            if text:
                ax.text(0.1, 0.9, text, fontsize=12)
            fig.savefig(str(path))
            plt.close(fig)

    def test_文本PDF抽取(self):
        with tempfile.TemporaryDirectory() as td:
            pdf = Path(td) / "d1.pdf"
            self._make_pdf(pdf, "Prior art D1 discloses a fixed silence window method for alerts.")
            r = pdf_to_text.extract(pdf)
            self.assertEqual(r["pages"], 1)
            self.assertIn("fixed silence window", r["text"])
            self.assertFalse(r["scanned_suspect"])

    def test_空页判扫描件(self):
        with tempfile.TemporaryDirectory() as td:
            pdf = Path(td) / "scan.pdf"
            self._make_pdf(pdf, None)
            r = pdf_to_text.extract(pdf)
            self.assertTrue(r["scanned_suspect"])

    def test_main_写出txt(self):
        with tempfile.TemporaryDirectory() as td:
            pdf = Path(td) / "x.pdf"
            self._make_pdf(pdf, "hello patent")
            rc = pdf_to_text.main([str(pdf), "--json"])
            self.assertEqual(rc, 0)
            self.assertIn("hello patent", (Path(td) / "x.txt").read_text(encoding="utf-8"))

    def _strip_glyph_names(self, src: Path, dst: Path):
        """把 Type3 字体 Differences 表里的字形名全改成 /uni00000000 这种"字体内部编号"。

        matplotlib 遇到没有字形名表的字体（macOS 的 Songti SC、很多 .ttc 中文字体）就是这么写的；
        pypdf 查不到 AGL 名字，只能把编号原样吐出来——这就是真实世界里"有字符没文字层"的 PDF。
        这样造出来的样本不看本机装了什么字体、不看 pypdf 新旧版本都成立。
        """
        from pypdf import PdfWriter
        from pypdf.generic import ArrayObject, NameObject
        writer = PdfWriter(clone_from=str(src))
        n = 0
        for page in writer.pages:
            for key, font in page["/Resources"]["/Font"].items():
                font = font.get_object()
                enc = font.get("/Encoding")
                diffs = enc.get_object().get("/Differences") if enc is not None else None
                if diffs is None:
                    continue
                new = ArrayObject()
                for item in diffs:
                    if isinstance(item, str):
                        new.append(NameObject(f"/uni{n:08X}"))
                        n += 1
                    else:
                        new.append(item)
                enc.get_object()[NameObject("/Differences")] = new
        writer.write(str(dst))
        return n

    def test_有字符但没文字层判乱码(self):
        """字体子集没带 ToUnicode 时，抽出来是 /uni0000004b 这种字形编号。

        这种页字符数不少，扫描件那条判据抓不到，但内容完全不能用——
        必须自己报出来，不能把编号当正文交给 agent。
        """
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        with tempfile.TemporaryDirectory() as td:
            plain = Path(td) / "plain.pdf"
            with plt.rc_context({"pdf.fonttype": 3}):      # Type3 才有 Differences 表
                self._make_pdf(plain, "Prior art D1 discloses a fixed silence window.")
            pdf = Path(td) / "garbled.pdf"
            renamed = self._strip_glyph_names(plain, pdf)
            self.assertGreater(renamed, 5, "样本 PDF 没造出 Type3 字形表，测试前提不成立")
            r = pdf_to_text.extract(pdf)
            self.assertFalse(r["scanned_suspect"], "字符数够多，不该判成扫描件")
            self.assertTrue(r["garbled_suspect"])
            self.assertGreaterEqual(r["garbled_ratio"], pdf_to_text.GARBLED_RATIO)
            self.assertIn("没有可用的文字层", r["text"])
            self.assertNotIn("fixed silence window", r["text"])

    def test_正常PDF不误报乱码(self):
        with tempfile.TemporaryDirectory() as td:
            pdf = Path(td) / "ok.pdf"
            self._make_pdf(pdf, "Prior art D1 discloses a fixed silence window method.")
            r = pdf_to_text.extract(pdf)
            self.assertFalse(r["garbled_suspect"])
            self.assertEqual(r["garbled_ratio"], 0.0)

    def test_乱码比例算法(self):
        self.assertEqual(pdf_to_text._garbled_ratio(""), 0.0)
        self.assertEqual(pdf_to_text._garbled_ratio("一种告警抑制方法"), 0.0)
        self.assertGreater(pdf_to_text._garbled_ratio("/uni0000004b/uni00000048"), 0.9)
        self.assertGreater(pdf_to_text._garbled_ratio("(cid:113)(cid:114)"), 0.9)
        self.assertGreater(pdf_to_text._garbled_ratio("/g23/g24/g25"), 0.9)


class EnvTest(unittest.TestCase):
    def test_结构完整(self):
        r = check_env.check(Path(tempfile.mkdtemp()))
        self.assertIn("python", r)
        self.assertEqual(set(r["modules"]),
                         {"docx", "pypdf", "matplotlib", "yaml", "latex2mathml", "pptx", "playwright"})
        self.assertEqual(r["modules"]["pptx"]["pip"], "python-pptx")
        self.assertEqual(set(r["hosts"]), {"Claude Code", "Codex CLI", "Cursor", "WorkBuddy", "nanobot"})
        self.assertFalse(any(h["present"] for h in r["hosts"].values()))

    def test_main_OK行(self):
        rc, out, _ = 跑main(check_env.main, [])
        self.assertEqual(rc, 0)
        self.assertTrue(最后一行(out).startswith("OK: python="), out)
        rc, out, err = 跑main(check_env.main, ["--json"])
        self.assertEqual(rc, 0)
        json.loads(out)                      # stdout 是干净 JSON
        self.assertTrue(最后一行(err).startswith("OK: "), err)


MATH_MD = """# 发明名称：一种带公式的方法

## 具体实施方式

匹配分 \\(s_{ij}\\) 由式 (1) 给出，其中 $\\rho \\in (0,1)$ 为平滑系数；价格 $5 和 $10 不是公式。

\\[ s_{ij} = b_{i} A_{j} + \\min\\left\\{1, \\frac{A_{j}}{\\max\\{\\varepsilon, b_{i}\\}}\\right\\} \\tag{1} \\]

$$
\\sigma = \\frac{1}{W} \\sum_{k=1}^{W} s_{(k)}
$$

| 符号 | 含义 |
|---|---|
| \\(|\\sigma - \\sigma'|\\) | 签名变化量 |
| $T_r$ | 最小重排间隔 |

- 判断 \\(\\Delta t \\geq T_r\\) 是否成立。
"""


class MathSpanTest(unittest.TestCase):
    """公式定位不依赖任何第三方库。"""

    def test_行内公式定位(self):
        spans = md_to_docx.inline_math_spans("a \\(x_1\\) b $y^2$ c")
        self.assertEqual([t[2] for t in spans], ["x_1", "y^2"])

    def test_美元不当公式(self):
        self.assertEqual(md_to_docx.inline_math_spans("价格 $5 和 $10"), [])
        self.assertEqual(md_to_docx.inline_math_spans("花了 $ 5 块"), [])
        self.assertEqual(md_to_docx.inline_math_spans("转义 \\$x\\$ 不算"), [])
        self.assertEqual([t[2] for t in md_to_docx.inline_math_spans("$x$")], ["x"])

    def test_块级公式识别(self):
        self.assertEqual(md_to_docx.block_math_start("\\[ a \\]"), ("\\[", "\\]"))
        self.assertEqual(md_to_docx.block_math_start("$$"), ("$$", "$$"))
        self.assertIsNone(md_to_docx.block_math_start("$x$ 开头的普通段"))

    def test_表格里公式的竖线不切列(self):
        cells = md_to_docx.split_table_row("| \\(|\\sigma - \\sigma'|\\) | 签名变化量 |")
        self.assertEqual(cells, ["\\(|\\sigma - \\sigma'|\\)", "签名变化量"])
        self.assertEqual(md_to_docx.split_table_row("| a | b |"), ["a", "b"])


@需要docx
@需要latex2mathml
class OmmlTest(unittest.TestCase):
    def _convert(self, td: Path, text: str = MATH_MD):
        (td / "f.md").write_text(text, encoding="utf-8")
        return md_to_docx.convert(td / "f.md", td / "f.docx")

    def test_公式转成Word原生公式(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            st = self._convert(td)
            # 行内 4（s_ij、rho、表格 2 个）+ 列表 1 + 块级 2 = 7
            self.assertEqual(st["omml"], 7)
            self.assertEqual(st["math_text"], 0)
            self.assertEqual(st["tables"], 1)
            self.assertEqual(st["headings"], 2)
            counts = math_to_omml.count_omml_in_docx(td / "f.docx")
            self.assertEqual(counts["oMath"], 7)
            self.assertEqual(counts["oMathPara"], 2)
            import zipfile
            xml = zipfile.ZipFile(td / "f.docx").read("word/document.xml").decode("utf-8")
            self.assertIn("Cambria Math", xml)
            self.assertIn('w:eastAsia="宋体"', xml)          # 正文字体没被公式带跑
            self.assertIn('w:eastAsia="黑体"', xml)
            self.assertIn("<m:nary>", xml)                    # ∑ 是真正的大算子
            self.assertIn('<m:chr m:val="∑"/>', xml)
            self.assertIn("<m:f>", xml)                       # 分式
            self.assertIn('<m:begChr m:val="{"/>', xml)      # \\left\\{ 定界符
            self.assertIn('<m:begChr m:val="|"/>', xml)      # |x| 绝对值配成定界符
            self.assertNotIn("\\(", xml)                         # 原文定界符没漏进正文
            footer = zipfile.ZipFile(td / "f.docx").read("word/footer1.xml").decode("utf-8")
            self.assertIn("PAGE", footer)                     # 页码还在
            # 美元金额原样保留
            import docx
            body = "\n".join(p.text for p in docx.Document(str(td / "f.docx")).paragraphs)
            self.assertIn("价格 $5 和 $10 不是公式", body)

    def test_转不了的公式原文保留并计数(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            st = self._convert(td, "段落 \\(\\frac{a}{\\) 坏公式，好公式 $x$。\n\n\\[ \\left( \\]\n")
            self.assertEqual(st["omml"], 1)
            self.assertEqual(st["math_text"], 2)
            import docx
            body = "\n".join(p.text for p in docx.Document(str(td / "f.docx")).paragraphs)
            self.assertIn("\\(\\frac{a}{\\)", body)
            self.assertIn("\\[ \\left( \\]", body)

    def test_没装latex2mathml时退化为原文(self):
        saved = {k: sys.modules.get(k) for k in ("latex2mathml", "latex2mathml.converter")}
        sys.modules["latex2mathml"] = None            # 让 import 抛 ImportError
        sys.modules["latex2mathml.converter"] = None
        try:
            self.assertFalse(math_to_omml.omml_available())
            self.assertIsNone(math_to_omml.try_latex_to_omml("x^2"))
            with tempfile.TemporaryDirectory() as td:
                td = Path(td)
                st = self._convert(td)
                self.assertEqual(st["omml"], 0)
                self.assertEqual(st["math_text"], 7)
                self.assertFalse(st["omml_available"])
                self.assertEqual(math_to_omml.count_omml_in_docx(td / "f.docx")["oMath"], 0)
                st2 = docx_to_md.convert(td / "f.docx", td / "back.md")
                back = (td / "back.md").read_text(encoding="utf-8")
                self.assertIn("\\(s_{ij}\\)", back)                 # 行内原文保留
                self.assertIn("\\sum_{k=1}^{W}", back)               # 块级原文保留
                self.assertIn("| \\(|\\sigma - \\sigma'|\\) | 签名变化量 |", back)
                self.assertEqual(st2["tables"], 1)
        finally:
            for k, v in saved.items():
                if v is None:
                    sys.modules.pop(k, None)
                else:
                    sys.modules[k] = v
        self.assertTrue(math_to_omml.omml_available())

    def test_main_统计行和OK行(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            (td / "f.md").write_text(MATH_MD, encoding="utf-8")
            rc, out, err = 跑main(md_to_docx.main, [str(td / "f.md")])
            self.assertEqual(rc, 0, err)
            self.assertIn("公式：omml=7 text=0", out)
            last = 最后一行(out)
            self.assertTrue(last.startswith("OK: "), out)
            self.assertIn("omml=7 text=0", last)
            self.assertIn("tables=1", last)


class MathToOmmlUnitTest(unittest.TestCase):
    def test_归一化(self):
        n = math_to_omml.normalize_latex_for_omml
        self.assertEqual(n(r"a \tag{2}"), r"a \quad (2)")
        self.assertEqual(n(r"x \le y"), r"x \leq y")
        self.assertEqual(n(r"\left( x \right)"), r"\left( x \right)")   # \left/\right 留给定界符
        self.assertEqual(n(r"\bigl( x \bigr)"), r"( x )")
        self.assertEqual(n(r"25^{\circ}C"), "25℃")
        self.assertEqual(n(""), "")

    @需要docx
    @需要latex2mathml
    def test_块级与行内元素类型(self):
        self.assertTrue(math_to_omml.latex_to_omml("x", display=True).tag.endswith("}oMathPara"))
        self.assertTrue(math_to_omml.latex_to_omml("x", display=False).tag.endswith("}oMath"))
        with self.assertRaises(ValueError):
            math_to_omml.latex_to_omml("   ")
        self.assertIsNone(math_to_omml.try_latex_to_omml(""))


@需要docx
class OkLineTest(unittest.TestCase):
    """全包统一的机读约定：成功最后一行 OK:，失败 stderr 一行 FAIL:。"""

    def test_docx_to_md(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            (td / "a.md").write_text("# 标题\n\n正文\n", encoding="utf-8")
            md_to_docx.convert(td / "a.md", td / "a.docx")
            rc, out, _ = 跑main(docx_to_md.main, [str(td / "a.docx")])
            self.assertEqual(rc, 0)
            self.assertTrue(最后一行(out).startswith("OK: paragraphs=1 headings=1"), out)
            rc, _, err = 跑main(docx_to_md.main, [str(td / "没有.docx")])
            self.assertEqual(rc, 1)
            self.assertTrue(err.startswith("FAIL: "), err)

    def test_md_to_docx_缺图FAIL(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            (td / "m.md").write_text("段落\n\n![](不存在.png)\n", encoding="utf-8")
            rc, out, err = 跑main(md_to_docx.main, [str(td / "m.md")])
            self.assertEqual(rc, 1)
            self.assertIn("FAIL: 缺图", err)
            self.assertNotIn("OK:", out)
            rc, _, err = 跑main(md_to_docx.main, [])
            self.assertEqual(rc, 1)                          # 参数错误也是 1，不和"缺依赖=2"混
            self.assertTrue(err.startswith("FAIL: "), err)

    @需要matplotlib
    def test_make_figure(self):
        with tempfile.TemporaryDirectory() as td:
            spec = Path(td) / "s.json"
            spec.write_text(json.dumps(FIG_SPEC), encoding="utf-8")
            rc, out, _ = 跑main(make_figure.main, [str(spec)])
            self.assertEqual(rc, 0)
            self.assertTrue(最后一行(out).startswith("OK: nodes=3 edges=2"), out)

    @需要pypdf
    @需要matplotlib
    def test_pdf_to_text_json时OK行去stderr(self):
        with tempfile.TemporaryDirectory() as td:
            pdf = Path(td) / "x.pdf"
            PdfTest._make_pdf(PdfTest(), pdf, "hello patent")
            rc, out, err = 跑main(pdf_to_text.main, [str(pdf), "--json"])
            self.assertEqual(rc, 0)
            json.loads(out)
            self.assertTrue(最后一行(err).startswith("OK: pages=1"), err)
            rc, out, _ = 跑main(pdf_to_text.main, [str(pdf)])
            self.assertTrue(最后一行(out).startswith("OK: pages=1"), out)
            self.assertIn("garbled=0", 最后一行(out))


if __name__ == "__main__":
    unittest.main()
